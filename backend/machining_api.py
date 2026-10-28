"""부품제작(가공 중 측정) — 황삭 · 중삭 · 정삭 단계 관리.

- 부품별 기준 템플릿(machining_templates.py)으로 각 OCR 항목을 어느 단계에서 재는지 정한다.
  같은 부품은 모양이 같고 치수만 바뀌므로 '도면 위 위치 + 표기 종류'로 항목을 맞춘다. 여유값은 라벨마다 고정.
- 앞 단계 측정이 모두 합격해야 다음 단계를 저장할 수 있다 (제품 번호별).
- 측정 전 측정기 0점(기준봉·기준 시편) 확인을 체크해야 저장된다.
- 제품 번호(시리얼)별로 기록하고, 자주검사 성적서(xlsx)로 내려받는다.
- 기존 measurement_records(조립·검증 화면용)와 따로 machining_stage_records 에 저장한다.
"""
import io
import json
import math
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Literal

from fastapi import APIRouter, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, ConfigDict, Field

from drawing_library import database
from machining_templates import MID_WINDOW, STAGE_KEYS, STAGE_NAMES, TEMPLATES
from measurement_api import current_snapshot, judge, measurement_state
from ocr_api import retrieval_query
from ocr_results import read_ocr_result
from part_identity import drawing_part
from purpose_errors.api import search_part
from purpose_errors.purposes import STORES, Purpose
from purpose_errors.vectors import query_active

router = APIRouter(prefix="/api/machining", tags=["부품제작 단계 측정"])

MATCH_DISTANCE = 0.08  # 도면 폭·높이 대비 위치 차이 허용값
SERIAL_RE = r"^[0-9A-Za-z가-힣._-]{1,40}$"
PLAN_LABEL = {
    "R": "황삭에서 마무리",
    "RF": "황삭·중삭 여유 확인 → 정삭 마무리",
    "F": "정삭 후 측정",
    "none": "가공 측정 안 함",
}


def init_machining():
    with database() as db:
        db.execute("""CREATE TABLE IF NOT EXISTS machining_stage_records (
            drawing_id TEXT NOT NULL REFERENCES drawings(drawing_id),
            ocr_revision TEXT NOT NULL,
            serial TEXT NOT NULL,
            item_id TEXT NOT NULL,
            stage TEXT NOT NULL CHECK(stage IN ('1','2','3')),
            measurement_mode TEXT NOT NULL,
            measured_value TEXT NOT NULL DEFAULT '',
            gauge_result TEXT,
            instrument TEXT NOT NULL DEFAULT '',
            note TEXT NOT NULL DEFAULT '',
            zero_checked INTEGER NOT NULL DEFAULT 0,
            judgement TEXT NOT NULL CHECK(judgement IN ('pass','fail','review')),
            judgement_reason TEXT NOT NULL,
            confirmed INTEGER NOT NULL DEFAULT 0,
            criteria_json TEXT NOT NULL,
            updated_at TEXT NOT NULL DEFAULT (datetime('now', 'localtime')),
            PRIMARY KEY (drawing_id, ocr_revision, serial, item_id, stage)
        )""")
        db.execute("""CREATE TABLE IF NOT EXISTS machining_stage_attempts (
            attempt_id INTEGER PRIMARY KEY AUTOINCREMENT,
            drawing_id TEXT NOT NULL,
            ocr_revision TEXT NOT NULL,
            serial TEXT NOT NULL,
            item_id TEXT NOT NULL,
            stage TEXT NOT NULL,
            record_json TEXT NOT NULL,
            created_at TEXT NOT NULL DEFAULT (datetime('now', 'localtime'))
        )""")


# --------------------------------------------------------------------------- 항목 ↔ 템플릿 맞추기
def signature(category: str, text: str) -> str:
    """표기 종류: 기하공차(G), 치수는 Ø/M/R/C/길이(L).

    기하공차 기호는 OCR이 비슷한 기호로 잘못 읽는 일이 있어(⌭↔↗) 기호로는 구분하지 않고 위치로만 맞춘다.
    """
    text = (text or "").strip()
    if category == "GDT":
        return "G"
    if text[:1].upper() in ("Ø", "⌀", "∅"):
        return "Ø"
    if re.match(r"^M\d", text):
        return "M"
    if re.match(r"^R\d", text):
        return "R"
    if re.match(r"^C\d", text):
        return "C"
    return "L"


def normalized(item: dict, document: dict):
    pos = item.get("position") or {}
    image = document.get("image") or {}
    width, height = image.get("width_px"), image.get("height_px")
    if not width or not height or "x" not in pos or "y" not in pos:
        return None
    return pos["x"] / width, pos["y"] / height


def match_template(document: dict, template: dict) -> dict:
    """item_id → 템플릿 항목. 가까운 쌍부터 하나씩 짝짓는다."""
    pairs = []
    for item in document.get("items", []):
        if item.get("excluded"):
            continue
        xy = normalized(item, document)
        if xy is None:
            continue
        sig = signature(item.get("category", ""), item.get("raw_text", ""))
        for entry in template["items"]:
            if entry["sig"] == sig:
                penalty = 0.0
            elif {entry["sig"], sig} == {"Ø", "L"}:
                penalty = 0.02  # OCR이 Ø 기호를 빠뜨린 경우도 위치가 같으면 맞춘다
            else:
                continue
            distance = math.dist(xy, entry["pos"]) + penalty
            if distance <= MATCH_DISTANCE:
                pairs.append((distance, item["id"], entry["key"]))
    pairs.sort()
    by_key = {entry["key"]: entry for entry in template["items"]}
    used_items, used_keys, result = set(), set(), {}
    for distance, item_id, key in pairs:
        if item_id in used_items or key in used_keys:
            continue
        used_items.add(item_id)
        used_keys.add(key)
        result[item_id] = {**by_key[key], "distance": round(distance, 4)}
    return result


def fmt(value) -> str:
    return f"{value:.3f}".rstrip("0").rstrip(".")


def stage_window(stage: str, allowance: float) -> float:
    if stage == "2":
        return MID_WINDOW
    return round(max(abs(allowance) / 4, 0.5), 3)


def stage_criteria(item: dict, stages: dict, stage: str) -> dict | None:
    """해당 단계에서 무엇으로 판정하는지. 그 단계에서 재지 않으면 None."""
    if stage not in stages:
        return None
    allowance = stages[stage]
    mode = item.get("measurement_mode") or "numeric"
    name = STAGE_NAMES[stage]
    if allowance:
        nominal = item.get("nominal_value")
        if mode != "numeric" or not isinstance(nominal, (int, float)):
            return {"role": "allowance", "mode": mode, "target": None, "lower_limit": None,
                    "upper_limit": None, "allowance": allowance,
                    "text": f"{name} 여유 확인 (기준치수 없음 — 관리자 확인 필요)"}
        window = stage_window(stage, allowance)
        target = nominal + allowance
        sign = "+" if allowance > 0 else "−"
        return {"role": "allowance", "mode": "numeric", "target": round(target, 6), "allowance": allowance,
                "lower_limit": round(target - window, 6), "upper_limit": round(target + window, 6),
                "text": f"{name} 목표 {fmt(target)} (도면 {fmt(nominal)} {sign} 여유 {fmt(abs(allowance))}, ±{fmt(window)})"}
    role = "final" if stage == "3" else "finish"
    return {"role": role, "mode": mode, "target": item.get("nominal_value"), "allowance": 0.0,
            "lower_limit": item.get("lower_limit"), "upper_limit": item.get("upper_limit"),
            "text": "도면 공차로 판정" if mode == "numeric" else (item.get("gauge_criteria") or "합격·불합격 기록")}


def part_info(drawing_id: str, document: dict):
    """가공 템플릿: 표시명이 아니라 부품 키(3D 연결 키 → 공통 별칭)로 고른다."""
    key = drawing_part(drawing_id, document)["key"]
    return key, TEMPLATES.get(key)


def build_plan(drawing_id: str, document: dict) -> dict:
    key, template = part_info(drawing_id, document)
    matched = match_template(document, template) if template else {}
    items = []
    for item in document.get("items", []):
        if item.get("excluded"):
            continue  # 관리자가 오인식으로 제외한 항목은 재지 않는다
        entry = matched.get(item["id"])
        if entry:
            plan, stages, label, source = entry["plan"], entry["stages"], entry["label"], "template"
        else:
            # 템플릿에 없는 항목은 마무리 후(정삭)에 잰다.
            plan, stages, label, source = "F", {"3": 0.0}, "", "rule"
        items.append({
            "item_id": item["id"], "raw_text": item.get("raw_text", ""), "plan": plan,
            "plan_label": PLAN_LABEL[plan], "label": label, "source": source,
            "template_key": entry["key"] if entry else None,
            "stages": {s: stage_criteria(item, stages, s) for s in STAGE_KEYS},
        })
    return {
        "part_key": key,
        "part_label": template["label"] if template else "",
        "stage_names": STAGE_NAMES,
        "stage_labels": template["stages"] if template else {s: f"{n} 후" for s, n in STAGE_NAMES.items()},
        "template_found": bool(template),
        "items": items,
    }


def read_records(drawing_id: str, revision: str, serial: str | None = None) -> list:
    query = """SELECT serial, item_id, stage, measurement_mode, measured_value, gauge_result, instrument, note,
                      zero_checked, judgement, judgement_reason, confirmed, criteria_json, updated_at
               FROM machining_stage_records WHERE drawing_id=? AND ocr_revision=?"""
    params = [drawing_id, revision]
    if serial is not None:
        query += " AND serial=?"
        params.append(serial)
    with database() as db:
        rows = db.execute(query, params).fetchall()
    records = []
    for row in rows:
        record = dict(row)
        record["confirmed"] = bool(record["confirmed"])
        record["zero_checked"] = bool(record["zero_checked"])
        record["status"] = measurement_state(record["judgement"], record["confirmed"])
        record["criteria"] = json.loads(record.pop("criteria_json"))
        record["work_purpose"] = Purpose.manufacturing.value
        records.append(record)
    return records


def stage_status(plan: dict, records: list) -> dict:
    done = {(r["item_id"], r["stage"]): r for r in records}
    status, unlocked = {}, True
    for stage in STAGE_KEYS:
        required = [i["item_id"] for i in plan["items"] if i["stages"][stage]]
        passed = [i for i in required if (done.get((i, stage)) or {}).get("status") == "completed"]
        failed = [i for i in required if (done.get((i, stage)) or {}).get("judgement") == "fail"]
        recorded = [i for i in required if (i, stage) in done]
        complete = len(passed) == len(required)
        status[stage] = {"required": len(required), "recorded": len(recorded), "passed": len(passed),
                         "failed": len(failed), "complete": bool(required) and complete, "unlocked": unlocked}
        unlocked = unlocked and complete  # 앞 단계를 다 합격해야 다음 단계가 열린다
    return status


def check_serial(serial: str) -> str:
    serial = serial.strip()
    if not re.fullmatch(SERIAL_RE, serial):
        raise HTTPException(422, "제품 번호는 1~40자의 영문·숫자·한글·-·_ 로 입력해주세요.")
    return serial


@router.get("/{drawing_id}/plan")
def get_plan(drawing_id: str, revision: str, serial: str = ""):
    snapshot = current_snapshot(drawing_id, revision)
    plan = build_plan(drawing_id, snapshot["result"])
    all_records = read_records(drawing_id, revision)
    serials = sorted({r["serial"] for r in all_records})
    serial = serial.strip()
    records = [r for r in all_records if r["serial"] == serial] if serial else []
    return {"drawing_id": drawing_id, "revision": revision, "serial": serial, "serials": serials,
            **plan, "records": records, "status": stage_status(plan, records)}


class StageInput(BaseModel):
    model_config = ConfigDict(extra="forbid")

    revision: str = Field(min_length=1, max_length=100)
    serial: str = Field(min_length=1, max_length=40)
    stage: Literal["1", "2", "3"]
    measured_value: str = Field(default="", max_length=100)
    gauge_result: Literal["pass", "fail"] | None = None
    confirmed: bool = False
    zero_checked: bool = False
    instrument: str = Field(default="", max_length=200)
    note: str = Field(default="", max_length=5000)


def parse_value(text: str) -> Decimal:
    try:
        if not re.fullmatch(r"[+-]?(?:[0-9]+(?:[.][0-9]*)?|[.][0-9]+)", text):
            raise InvalidOperation
        return Decimal(text)
    except InvalidOperation:
        raise HTTPException(422, "실측값에는 유효한 숫자를 입력해주세요.")


def judge_allowance(criteria: dict, number: Decimal, stage_name: str):
    low, high = criteria["lower_limit"], criteria["upper_limit"]
    if low is None or high is None:
        return "review", "목표 치수를 계산할 수 없습니다. 관리자 확인이 필요합니다."
    if Decimal(str(low)) <= number <= Decimal(str(high)):
        return "pass", f"{stage_name}: 여유가 목표 범위 안입니다."
    short = number < Decimal(str(low))
    hole = criteria["allowance"] < 0
    # 축은 작으면, 구멍·홈은 크면 여유 부족(과절삭)
    if short != hole:
        return "fail", f"{stage_name}: 여유가 부족합니다(과절삭). 다음 단계에서 도면 치수를 만들 수 있는지 확인하세요."
    return "fail", f"{stage_name}: 여유가 너무 많습니다. {stage_name} 추가 가공이 필요합니다."


@router.post("/{drawing_id}/items/{item_id}/record")
def save_stage_record(drawing_id: str, item_id: str, payload: StageInput):
    snapshot = current_snapshot(drawing_id, payload.revision)
    document = snapshot["result"]
    if document.get("review_status") != "reviewed":
        raise HTTPException(409, "관리자 검토 완료 후 측정 결과를 저장할 수 있습니다.")
    serial = check_serial(payload.serial)
    stage, stage_name = payload.stage, STAGE_NAMES[payload.stage]

    plan = build_plan(drawing_id, document)
    item = next((entry for entry in document["items"] if entry["id"] == item_id), None)
    if item is None:
        raise HTTPException(404, "이 도면의 항목이 아닙니다.")
    entry = next(i for i in plan["items"] if i["item_id"] == item_id)
    criteria = entry["stages"][stage]
    if criteria is None:
        raise HTTPException(422, f"이 항목은 {stage_name} 측정 대상이 아닙니다.")

    status = stage_status(plan, read_records(drawing_id, payload.revision, serial))
    if not status[stage]["unlocked"]:
        previous = STAGE_KEYS[STAGE_KEYS.index(stage) - 1]
        s = status[previous]
        raise HTTPException(
            409,
            f"{STAGE_NAMES[previous]} 측정이 모두 합격해야 {stage_name}을 기록할 수 있습니다. "
            f"({STAGE_NAMES[previous]} 합격 {s['passed']}/{s['required']})",
        )
    if not payload.zero_checked:
        raise HTTPException(422, "측정 전에 기준봉·기준 시편으로 측정기 0점을 확인하고 체크해주세요.")

    mode = criteria["mode"]
    value = payload.measured_value.strip()
    if mode == "gauge":
        if value or payload.gauge_result is None:
            raise HTTPException(422, "합격·불합격 기록 항목입니다. 결과를 선택해주세요.")
        judgement = payload.gauge_result
        reason = f"{stage_name}: 작업자가 판정 기준에 따라 " + ("적합으로 기록했습니다." if judgement == "pass"
                                                          else "부적합으로 기록했습니다. 재확인이 필요합니다.")
    else:
        if payload.gauge_result is not None:
            raise HTTPException(422, "실측값 방식에서는 게이지 결과를 보낼 수 없습니다.")
        number = parse_value(value)
        if criteria["role"] == "allowance":
            judgement, reason = judge_allowance(criteria, number, stage_name)
        else:
            judgement, reason = judge(item, number)
            if judgement != "review":
                reason = f"{stage_name}: {reason}"

    if payload.confirmed and judgement != "pass":
        raise HTTPException(422, "적합 판정인 항목만 측정 완료로 확인할 수 있습니다.")

    record = {
        "drawing_id": drawing_id, "revision": payload.revision, "serial": serial, "item_id": item_id,
        "stage": stage, "measurement_mode": mode, "measured_value": value if mode == "numeric" else "",
        "gauge_result": payload.gauge_result, "instrument": payload.instrument.strip(),
        "note": payload.note.strip(), "zero_checked": True, "judgement": judgement,
        "judgement_reason": reason, "confirmed": payload.confirmed,
        "status": measurement_state(judgement, payload.confirmed), "criteria": criteria,
        "work_purpose": Purpose.manufacturing.value,
    }
    with database() as db:
        db.execute("BEGIN IMMEDIATE")
        current = db.execute("SELECT revision FROM ocr_results WHERE drawing_id=?", (drawing_id,)).fetchone()
        if current is None or current["revision"] != payload.revision:
            raise HTTPException(409, "저장 중 OCR 결과가 변경됐습니다. 다시 조회해주세요.")
        db.execute(
            """INSERT INTO machining_stage_records (drawing_id, ocr_revision, serial, item_id, stage,
                   measurement_mode, measured_value, gauge_result, instrument, note, zero_checked, judgement,
                   judgement_reason, confirmed, criteria_json)
               VALUES (?,?,?,?,?,?,?,?,?,?,1,?,?,?,?)
               ON CONFLICT(drawing_id, ocr_revision, serial, item_id, stage) DO UPDATE SET
                   measurement_mode=excluded.measurement_mode, measured_value=excluded.measured_value,
                   gauge_result=excluded.gauge_result, instrument=excluded.instrument, note=excluded.note,
                   zero_checked=1, judgement=excluded.judgement, judgement_reason=excluded.judgement_reason,
                   confirmed=excluded.confirmed, criteria_json=excluded.criteria_json,
                   updated_at=datetime('now', 'localtime')""",
            (drawing_id, payload.revision, serial, item_id, stage, mode, record["measured_value"],
             payload.gauge_result, record["instrument"], record["note"], judgement, reason,
             int(payload.confirmed), json.dumps(criteria, ensure_ascii=False)),
        )
        # 재측정해도 이전 결과는 이력으로 남긴다.
        db.execute(
            """INSERT INTO machining_stage_attempts (drawing_id, ocr_revision, serial, item_id, stage, record_json)
               VALUES (?,?,?,?,?,?)""",
            (drawing_id, payload.revision, serial, item_id, stage, json.dumps(record, ensure_ascii=False)),
        )
    return record


# --------------------------------------------------------------------------- 단계별 오류사례 검색
@router.get("/{drawing_id}/items/{item_id}/errors")
def stage_errors(drawing_id: str, item_id: str, request: Request, stage: Literal["1", "2", "3"],
                 top_k: int = Query(default=3, ge=1, le=10)):
    """부품제작(manufacturing) 오류사례 중 선택한 단계(황삭 후·중삭 후·정삭 후)와 '가공 공통'만 찾는다.

    공정을 비워 둔 사용자 등록 사례는 모든 단계에 보인다.
    """
    snapshot = read_ocr_result(drawing_id)
    document = snapshot["result"]
    if document is None:
        raise HTTPException(409, "OCR 결과를 먼저 저장해주세요.")
    item = next((entry for entry in document["items"] if entry["id"] == item_id), None)
    if item is None:
        raise HTTPException(404, "이 도면의 측정항목을 찾을 수 없습니다.")
    part = search_part(STORES[Purpose.manufacturing], drawing_id, snapshot)
    families = part["families"]
    if not families:
        raise HTTPException(409, "도면의 소속 부품을 확인해주세요.")

    collections = getattr(request.app.state, "purpose_collections", None) or {}
    collection = collections.get(Purpose.manufacturing)
    if collection is None:
        raise HTTPException(503, "부품제작 오류사례 검색 저장소가 준비되지 않았습니다.")

    conditions = [{"part_family": {"$in": families}},
                  {"process_stage": {"$in": [f"{STAGE_NAMES[stage]} 후", "가공 공통", ""]}}]
    feature = item.get("feature_type")
    if feature and feature != "UNKNOWN":
        conditions.append({"feature_type": feature})
    where = {"$and": conditions}
    candidates = collection.get(where=where, include=[])
    errors = []
    if candidates["ids"]:
        # 비활성 사례는 원문 테이블 기준으로 뺀다 (목적별 검색과 같은 규칙).
        errors = query_active(STORES[Purpose.manufacturing], collection, request.app.state.model,
                              query_text=retrieval_query(item), where=where,
                              candidate_ids=candidates["ids"], top_k=top_k)
    return {"purpose": Purpose.manufacturing.value, "drawing_id": drawing_id, "item_id": item_id,
            "stage": stage, "part": part, "errors": errors}


# --------------------------------------------------------------------------- 자주검사 성적서
@router.get("/{drawing_id}/report.xlsx")
def stage_report(drawing_id: str, revision: str):
    """치수 검사 기록서(자주검사) — 행: 항목×단계, 열: 제품 번호별 실측·편차·판정."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    snapshot = current_snapshot(drawing_id, revision)
    document = snapshot["result"]
    plan = build_plan(drawing_id, document)
    records = read_records(drawing_id, revision)
    serials = sorted({r["serial"] for r in records})
    by_key = {(r["serial"], r["item_id"], r["stage"]): r for r in records}
    with database() as db:
        row = db.execute("SELECT original_filename FROM drawings WHERE drawing_id=?", (drawing_id,)).fetchone()
    drawing_name = row["original_filename"] if row else drawing_id
    number_of = {item["id"]: index + 1 for index, item in enumerate(document.get("items", []))}

    wb = Workbook()
    ws = wb.active
    ws.title = "치수검사기록서"
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="17365D")
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.append(["치수 검사 기록서 (자주검사 · 가공 중 측정)"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["부품", plan["part_label"] or document.get("part_name") or "", "도면", drawing_name])
    ws.append(["OCR 버전", revision, "출력", datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws.append(["판정: ○ 합격 · ✕ 불합격 · △ 확인 필요 / 편차 = 실측 − 목표(목표가 없으면 도면 기준치수)"])
    ws.append([])
    header = ["CP", "항목", "도면 표기", "단계", "판정 기준"]
    for serial in serials:
        header += [f"{serial} 실측", f"{serial} 편차", f"{serial} 판정"]
    ws.append(header)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font, cell.fill, cell.border = head, fill, box
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    mark = {"pass": "○", "fail": "✕", "review": "△"}
    rows = []
    for entry in plan["items"]:
        for stage in STAGE_KEYS:
            criteria = entry["stages"][stage]
            if criteria:
                rows.append((f"{number_of.get(entry['item_id'], '')}", entry["label"] or entry["item_id"],
                             entry["raw_text"].replace("\n", " "), stage, criteria, entry["item_id"]))
    rows.sort(key=lambda r: (r[3], int(r[0]) if r[0].isdigit() else 999))
    for cp, label, text, stage, criteria, item_id in rows:
        basis = criteria["text"]
        if criteria["lower_limit"] is not None and criteria["upper_limit"] is not None and criteria["mode"] == "numeric":
            basis += f" [{fmt(criteria['lower_limit'])} ~ {fmt(criteria['upper_limit'])}]"
        line = [cp, label, text, STAGE_NAMES[stage], basis]
        for serial in serials:
            record = by_key.get((serial, item_id, stage))
            if not record:
                line += ["", "", ""]
                continue
            reference = criteria["target"]
            deviation = ""
            if record["measured_value"] and isinstance(reference, (int, float)):
                deviation = round(float(record["measured_value"]) - reference, 4)
            measured = record["measured_value"] or ("합격" if record["gauge_result"] == "pass" else "불합격")
            line += [measured, deviation, mark[record["judgement"]]]
        ws.append(line)
        for cell in ws[ws.max_row]:
            cell.border = box
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [6, 22, 16, 7, 42] + [10, 9, 6] * len(serials)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = ws.cell(header_row + 1, 6)

    log = wb.create_sheet("측정 이력")
    log.append(["제품 번호", "CP", "항목", "단계", "실측", "판정", "판정 사유", "측정기", "0점 확인", "메모", "기록 시각"])
    for cell in log[1]:
        cell.font, cell.fill = head, fill
    labels = {i["item_id"]: (i["label"] or i["item_id"]) for i in plan["items"]}
    for record in sorted(records, key=lambda r: (r["serial"], r["stage"], r["updated_at"])):
        log.append([record["serial"], number_of.get(record["item_id"], ""), labels.get(record["item_id"], record["item_id"]),
                    STAGE_NAMES[record["stage"]],
                    record["measured_value"] or record["gauge_result"], mark[record["judgement"]],
                    record["judgement_reason"], record["instrument"], "확인" if record["zero_checked"] else "",
                    record["note"], record["updated_at"]])
    for index, width in enumerate([10, 6, 22, 7, 10, 6, 46, 18, 8, 20, 18], start=1):
        log.column_dimensions[get_column_letter(index)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    stem = re.sub(r"[^0-9A-Za-z_-]+", "_", drawing_name.rsplit(".", 1)[0])
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{stem}_machining_report.xlsx"'},
    )
